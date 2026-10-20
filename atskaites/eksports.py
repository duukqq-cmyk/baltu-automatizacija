"""Atskaites eksports uz Excel (openpyxl)."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from atskaites.metrikas import Atskaite
from atskaites.modelis import LAUKI, POSMU_NOSAUKUMI, Dati

_BOLD = Font(bold=True)
_HEAD = PatternFill("solid", fgColor="DDEBF7")
_SADALA = PatternFill("solid", fgColor="F2F2F2")


def _tabula(ws, galvas: list[str], rindas: list[list], platumi: list[int] | None = None, sakums: int = 1):
    for c, g in enumerate(galvas, start=1):
        cell = ws.cell(sakums, c, g)
        cell.font = _BOLD
        cell.fill = _HEAD
    for r, rinda in enumerate(rindas, start=sakums + 1):
        for c, v in enumerate(rinda, start=1):
            cell = ws.cell(r, c, v)
            if hasattr(v, "year"):
                cell.number_format = "DD.MM.YYYY"
    for i, w in enumerate(platumi or [], start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = ws.cell(sakums + 1, 1)


def excel_atskaite(cels: str, a: Atskaite, dati: Dati) -> None:
    wb = Workbook()

    ws = wb.active
    ws.title = "Kopsavilkums"
    ws.cell(1, 1, f"Leadu atskaite — {a.periods.nosaukums}").font = Font(bold=True, size=14)
    ws.cell(2, 1, f"Periods: {a.periods.no:%d.%m.%Y} – {a.periods.lidz:%d.%m.%Y}")
    r = 4
    for c, g in enumerate(["Rādītājs", "Vērtība", "Vienība", "Piezīme"], start=1):
        ws.cell(r, c, g).font = _BOLD
        ws.cell(r, c).fill = _HEAD
    sadala = None
    for m in a.metriki:
        if m.sadala != sadala:
            sadala = m.sadala
            r += 1
            ws.cell(r, 1, sadala).font = _BOLD
            for c in range(1, 5):
                ws.cell(r, c).fill = _SADALA
        r += 1
        ws.cell(r, 1, m.nosaukums)
        ws.cell(r, 2, "—" if m.vertiba is None else m.vertiba)
        ws.cell(r, 3, m.vieniba)
        ws.cell(r, 4, m.piezime).alignment = Alignment(wrap_text=True)
    if a.atteikuma_iemesli:
        r += 2
        ws.cell(r, 1, "Atteikuma iemesli").font = _BOLD
        for iem, n in a.atteikuma_iemesli:
            r += 1
            ws.cell(r, 1, iem)
            ws.cell(r, 2, n)
    for kol, w in zip("ABCD", [48, 14, 12, 80]):
        ws.column_dimensions[kol].width = w

    if a.pa_agentiem:
        ws = wb.create_sheet("Pa agentiem")
        galvas = list(a.pa_agentiem[0].keys())
        _tabula(ws, galvas, [[r[g] for g in galvas] for r in a.pa_agentiem], [26] + [13] * (len(galvas) - 1))

    ws = wb.create_sheet("Piltuve tagad")
    _tabula(ws, ["Posms", "Leadi"], [[p["Posms"], p["Leadi"]] for p in a.piltuve_tagad], [30, 10])

    ws = wb.create_sheet("Leadi")
    if a.leadu_tabula:
        galvas = list(a.leadu_tabula[0].keys())
        _tabula(ws, galvas, [[r[g] for g in galvas] for r in a.leadu_tabula],
                [34, 14, 13, 24, 7, 12, 12, 22, 11, 7, 34, 10, 9, 11] + [12] * 7 + [22, 60])

    ws = wb.create_sheet("Datu kvalitāte")
    _tabula(ws, ["Agents", "Lead", "Problēma"],
            [[p["Agents"], p["Lead"], p["Problēma"]] for p in a.datu_kvalitate], [26, 34, 80])

    ws = wb.create_sheet("Kolonnu karte")
    galvas = ["Standarta lauks", "Tips"] + [b.nosaukums for b in dati.boardi]
    rindas = []
    for l in LAUKI:
        rindas.append([l.nosaukums, l.tips] + [dati.kartes.get(b.id, {}).get(l.key, "— nav —") for b in dati.boardi])
    _tabula(ws, galvas, rindas, [32, 8] + [28] * len(dati.boardi))
    r = len(rindas) + 3
    ws.cell(r, 1, "Grupa monday").font = _BOLD
    ws.cell(r, 2, "Posms").font = _BOLD
    for g, p in dati.grupu_karte.items():
        r += 1
        ws.cell(r, 1, g)
        ws.cell(r, 2, POSMU_NOSAUKUMI.get(p, p))

    wb.save(cels)
