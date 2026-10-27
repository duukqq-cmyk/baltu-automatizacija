"""Gala atskaites ģenerēšana (teksts + Excel)."""
from __future__ import annotations

from .formula import SUGU_NOSAUKUMI
from .parser import Inventarizacija
from .rules import Vertejums
from .sortiment import FinansuRezultats


def _pdf_info(inv: Inventarizacija, n) -> list[str]:
    """gabali.lv papildu informācija zem nogabala rindas (VMD — nav)."""
    return n.pdf_apraksts if inv.avots == "gabali.lv" else []


def _rezerves_proc(fin: FinansuRezultats) -> float:
    segums = fin.max_iegades_summa + fin.pelna
    return fin.pelna / segums if segums else 0.0


def nog_nosaukums(inv: Inventarizacija, n) -> str:
    """Nogabala nosaukums atskaitē; kvartālu rāda tikai, ja īpašumā to ir vairāki."""
    return f"{n.apzimejums}.nog" if inv.vairaki_kvartali else f"{n.nr}.nog"


NAV_MEZAUDZE = "Nav mežaudze"
NECERTAMS = "Netiek cirsts"
KAILCIRTE = "Kailcirte"


def _vmd_cirtes_veids(planots: str | None) -> str | None:
    """'Koku ciršana 2026, Kailcirte' -> 'Kailcirte'."""
    if not planots:
        return None
    veids = planots.rsplit(",", 1)[-1].strip()
    return veids[:1].upper() + veids[1:] if veids else None


def platiba_pa_cirtes_veidiem(inv: Inventarizacija, vertejumi: list[Vertejums]) -> list[dict]:
    """Platības kopsavilkums pēc cirtes veida.

    Kailcirte — nogabali, kas atbilst kailcirtei (aprēķinā iekļautie).
    Pārējām mežaudzēm — VMD plānotās cirtes veids (ja avotā norādīts),
    citādi "Netiek cirsts". Zeme, kas nav mežaudze, — atsevišķā rindā.
    Atgriež rindas {"Cirtes veids", "Nogabali", "Platība ha", "Daļa"} + "Kopā".
    """
    grupas: dict[str, list] = {}
    indeksi = {v.indekss for v in vertejumi}
    objekti = {id(v.nogabals) for v in vertejumi}
    for v in vertejumi:
        if v.atlauta_kailcirte:
            veids = KAILCIRTE
        else:
            veids = _vmd_cirtes_veids(v.nogabals.planots)
            veids = f"{veids} (VMD plānots)" if veids else NECERTAMS
        grupas.setdefault(veids, []).append(v.nogabals)
    for i, n in enumerate(inv.nogabali):
        if i not in indeksi and id(n) not in objekti:  # nav novērtēts -> nav mežaudze
            grupas.setdefault(NAV_MEZAUDZE, []).append(n)

    def kartiba(veids: str) -> tuple:
        return (veids != KAILCIRTE, veids in (NECERTAMS, NAV_MEZAUDZE), veids == NAV_MEZAUDZE, veids)

    kopa = sum(n.platiba_ha or 0 for ng in grupas.values() for n in ng)
    rindas = []
    for veids in sorted(grupas, key=kartiba):
        ha = round(sum(n.platiba_ha or 0 for n in grupas[veids]), 2)
        rindas.append({
            "Cirtes veids": veids,
            "Nogabali": ";".join(nog_nosaukums(inv, n).removesuffix(".nog") for n in grupas[veids]),
            "Platība ha": ha,
            "Daļa": ha / kopa if kopa else 0.0,
        })
    rindas.append({"Cirtes veids": "Kopā", "Nogabali": "",
                   "Platība ha": round(kopa, 2), "Daļa": 1.0 if kopa else 0.0})
    return rindas


def atrast_apjomu(fin: FinansuRezultats, n):
    """Nogabala apjoma ieraksts (pēc objekta — nogabalu numuri kvartālos atkārtojas)."""
    return next((a for a in fin.apjomi if a.nogabals is n), None)


def teksta_atskaite(inv: Inventarizacija, vertejumi: list[Vertejums], fin: FinansuRezultats) -> str:
    """Gala atskaite pamācībā prasītajā formātā + izvērsts pamatojums."""
    cirt = [v for v in vertejumi if v.atlauta_kailcirte]
    necirt = [v for v in vertejumi if not v.atlauta_kailcirte]

    nog_saraksts = ";".join(nog_nosaukums(inv, v.nogabals) for v in cirt)
    sugas = "; ".join(f"{s}={round(m3)}m3" for s, m3 in fin.pa_sugam_m3.items())
    summa_txt = f"{fin.max_iegades_summa:,.0f}".replace(",", " ")

    rindas = [
        "CIRSMAS VĒRTĪBAS APRĒĶINS",
        "=" * 60,
        "",
        f"NĪ {inv.ipasums or '—'} , zemes vienība {inv.zemes_vieniba or '—'}",
        f"Datu avots: {inv.avots}",
        "",
        f"Kailcirtes: {nog_saraksts}",
        "",
        f"Cērtamais apjoms: {round(fin.kopejais_m3)}m3 ( {sugas} )",
        "",
        f"Maksimālā iegādes summa = {summa_txt} EUR, jeb {fin.cena_eur_m3:.2f} EUR/m3.",
        "",
        "-" * 60,
        "DETALIZĒTS APRĒĶINS",
        "-" * 60,
        "",
        "Kailcirtei atbilstošie nogabali:",
    ]
    for v in cirt:
        n = v.nogabals
        apjoms = atrast_apjomu(fin, n)
        kraja = f"{n.kraja_m3_ha:.0f} m3/ha" if n.kraja_m3_ha is not None else "krāja pēc sugām"
        metode = f" [{apjoms.metode}]" if apjoms and apjoms.metode != "formula" else ""
        rindas.append(
            f"  {nog_nosaukums(inv, n)} | {n.platiba_ha} ha x {kraja} = "
            f"{apjoms.kopejais_m3 if apjoms else 0:.2f} m3{metode} | {n.formula} | {v.pamatojums}"
        )
        for b in v.bridinajumi:
            rindas.append(f"      {b}")
        for r in _pdf_info(inv, n):
            rindas.append(f"      • {r}")
    rindas += ["", "Kailcirtei NEatbilstošie nogabali:"]
    for v in necirt:
        rindas.append(f"  {nog_nosaukums(inv, v.nogabals)} | {v.pamatojums}")
        for b in v.bridinajumi:
            if b.startswith("✎"):
                rindas.append(f"      {b}")
        for r in _pdf_info(inv, v.nogabals):
            rindas.append(f"      • {r}")

    riski = [v for v in cirt if any(b.startswith("⚠") for b in v.bridinajumi)]
    if riski:
        rindas += ["", f"⚠ UZMANĪBU: {len(riski)} no {len(cirt)} cērtamajiem nogabaliem ir riska "
                   "karogi (aizsargjoslas, VMD plānotā darbība vai kontrolsummas) — skat. augstāk."]

    rindas += ["", "Apjoms pa sugām:"]
    for s, m3 in fin.pa_sugam_m3.items():
        rindas.append(f"  {SUGU_NOSAUKUMI.get(s, s):14s} {m3:10.2f} m3")

    rindas += ["", "Platība pēc cirtes veidiem:"]
    platibas = platiba_pa_cirtes_veidiem(inv, vertejumi)
    platums = max(27, *(len(r["Cirtes veids"]) + 2 for r in platibas))
    for r in platibas:
        nog = f"  ({r['Nogabali']})" if r["Nogabali"] else ""
        rindas.append(f"  {r['Cirtes veids'] + ':':{platums}s}{r['Platība ha']:8.2f} ha {r['Daļa']:6.1%}{nog}")

    rindas += [
        "",
        "Finanses:",
        f"  Ieņēmumi no sortimentiem:  {fin.ienemumi_kopa:12,.2f} EUR".replace(",", " "),
    ]
    for poz, eur in fin.izdevumi.items():
        rindas.append(f"  {poz + ':':27s}{-eur:12,.2f} EUR".replace(",", " "))
    rindas += [
        f"  {'Peļņas rezerve (' + format(_rezerves_proc(fin), '.0%') + '):':27s}{-fin.pelna:12,.2f} EUR".replace(",", " "),
        f"  MAX IEGĀDES SUMMA:         {fin.max_iegades_summa:12,.2f} EUR".replace(",", " "),
        f"  Cena:                      {fin.cena_eur_m3:12.2f} EUR/m3",
    ]
    return "\n".join(rindas)


def excel_atskaite(cels: str, inv: Inventarizacija, vertejumi: list[Vertejums], fin: FinansuRezultats) -> None:
    """Izveido Excel atskaiti ar 3 lapām: kopsavilkums, nogabali, sortimenti."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D9E1F2")

    # 1. lapa — kopsavilkums
    ws = wb.active
    ws.title = "Kopsavilkums"
    cirt = [v for v in vertejumi if v.atlauta_kailcirte]
    dati = [
        ("Objekts", f"NĪ {inv.ipasums or '—'}, zemes vienība {inv.zemes_vieniba or '—'}"),
        ("Datu avots", inv.avots),
        ("Kailcirtes", ";".join(nog_nosaukums(inv, v.nogabals) for v in cirt)),
        ("Nogabali ar riska karogiem (⚠)",
         sum(1 for v in cirt if any(b.startswith("⚠") for b in v.bridinajumi))),
        ("Cērtamais apjoms, m3", fin.kopejais_m3),
        *[(f"  t.sk. {SUGU_NOSAUKUMI.get(s, s)} ({s}), m3", m3) for s, m3 in fin.pa_sugam_m3.items()],
        *[(f"Platība: {r['Cirtes veids']}, ha" if r["Cirtes veids"] != "Kopā" else "Platība kopā, ha",
           r["Platība ha"]) for r in platiba_pa_cirtes_veidiem(inv, vertejumi)],
        ("Ieņēmumi kopā, EUR", fin.ienemumi_kopa),
        *[(f"Izdevumi: {p}, EUR", -e) for p, e in fin.izdevumi.items()],
        ("Peļņas rezerve, EUR", -fin.pelna),
        ("MAKSIMĀLĀ IEGĀDES SUMMA, EUR", fin.max_iegades_summa),
        ("Cena, EUR/m3", fin.cena_eur_m3),
    ]
    for r, (k, val) in enumerate(dati, start=1):
        ws.cell(r, 1, k).font = bold
        ws.cell(r, 2, val)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 46

    # 2. lapa — nogabalu vērtējums
    ws2 = wb.create_sheet("Nogabali")
    galvas = ["Nog.", "Platība ha", "Krāja m3/ha", "Formula", "Bonitāte", "Caurmērs cm",
              "Vecums", "Valdošā suga", "Kailcirte?", "Pamatojums", "Apjoms m3",
              "Apjoma metode", "Apjoms pa sugām", "VMD plānots", "Brīdinājumi", "Info no PDF"]
    for c, g in enumerate(galvas, start=1):
        cell = ws2.cell(1, c, g)
        cell.font = bold
        cell.fill = header_fill
    for r, v in enumerate(vertejumi, start=2):
        n = v.nogabals
        a = atrast_apjomu(fin, n)
        sugas = "; ".join(f"{s_}={m3:.1f}" for s_, m3 in a.pa_sugam.items()) if a else ""
        for c, val in enumerate([nog_nosaukums(inv, n), n.platiba_ha, n.kraja_m3_ha, n.formula,
                                 n.bonitate, n.vid_caurmers, n.vecums, v.valdosa_suga,
                                 "JĀ" if v.atlauta_kailcirte else "NĒ", v.pamatojums,
                                 a.kopejais_m3 if a else "", a.metode if a else "", sugas,
                                 n.planots or "", "\n".join(v.bridinajumi),
                                 "\n".join(_pdf_info(inv, n))],
                                start=1):
            ws2.cell(r, c, val)
    for kol, w in zip("ABCDEFGHIJKLMNOP", [12, 10, 11, 22, 9, 12, 8, 12, 10, 60, 11, 13, 30, 30, 80, 80]):
        ws2.column_dimensions[kol].width = w

    # 3. lapa — sortimentu aprēķins
    ws3 = wb.create_sheet("Sortimenti")
    galvas3 = ["Suga", "Sortiments", "Īpatsvars", "Apjoms m3", "Cena EUR/m3", "Summa EUR"]
    for c, g in enumerate(galvas3, start=1):
        cell = ws3.cell(1, c, g)
        cell.font = bold
        cell.fill = header_fill
    r = 2
    for rinda in fin.detalas:
        for c, k in enumerate(galvas3, start=1):
            ws3.cell(r, c, rinda[k])
        r += 1
    ws3.cell(r, 5, "KOPĀ:").font = bold
    ws3.cell(r, 6, fin.ienemumi_kopa).font = bold
    for kol, w in zip("ABCDEF", [10, 14, 10, 11, 12, 12]):
        ws3.column_dimensions[kol].width = w

    wb.save(cels)
